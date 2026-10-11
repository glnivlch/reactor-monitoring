# Библиотеки
from datetime import datetime
from openpyxl import load_workbook

# Константы
TEMPERATURE_MIN = 70.0
TEMPERATURE_MAX = 90.0
PRESSURE_MIN = 2.0
PRESSURE_MAX = 5.0
history = []


# Функции
# Проверка параметра на допустимый диапазон
def check_parameter(
    parameter: float,
    lower_limit: float,
    upper_limit: float
) -> str:
    if lower_limit <= parameter <= upper_limit:
        return "normal"
    elif parameter > upper_limit:
        return "high"
    else:
        return "low"


# Сохранение параметров в историю измерений
def add_measurement(temperature: float, pressure: float) -> tuple[str, str]:
    measurement_time = datetime.now()
    measurement = (measurement_time, temperature, pressure)
    history.append(measurement)
    save_to_excel(measurement)

    temperature_status = check_parameter(temperature, TEMPERATURE_MIN, TEMPERATURE_MAX)
    pressure_status = check_parameter(pressure, PRESSURE_MIN, PRESSURE_MAX)

    return (temperature_status, pressure_status)


# Сохранение измерения в Excel
def save_to_excel(measurement: tuple) -> None:
    book = load_workbook("measurements.xlsx")
    sheet = book.active
    sheet.append(measurement)
    book.save("measurements.xlsx")


# Восстановление предыдущей сессии в актуальную историю измерений
def load_history_from_excel() -> None:
    book = load_workbook("measurements.xlsx")
    sheet = book.active
    for i in range(2, sheet.max_row + 1):
        measurement = []
        for j in range(1, 4):
            measurement.append(sheet.cell(row=i, column=j).value)

        measurement = tuple(measurement)
        history.append(measurement)


# Расчет статистики параметров
def statistics() -> tuple[float, float, float, float, float, float]:
    # Фиксация первых измерений, как экстремумов
    max_temperature = history[0][1]
    min_temperature = history[0][1]
    max_pressure = history[0][2]
    min_pressure = history[0][2]

    total_temperature = 0
    total_pressure = 0

    # Поиск экстремумов параметров среди всех измерений
    for measurement_saved in history:
        if measurement_saved[1] > max_temperature:
            max_temperature = measurement_saved[1]
        if measurement_saved[1] < min_temperature:
            min_temperature = measurement_saved[1]

        if measurement_saved[2] > max_pressure:
            max_pressure = measurement_saved[2]
        if measurement_saved[2] < min_pressure:
            min_pressure = measurement_saved[2]

        total_temperature += measurement_saved[1]
        total_pressure += measurement_saved[2]

    average_temperature = total_temperature / len(history)
    average_pressure = total_pressure / len(history)

    return (max_temperature, min_temperature, average_temperature,
            max_pressure, min_pressure, average_pressure)
